import os, io, zipfile, tempfile, shutil, json, threading, time, uuid, traceback
from pathlib import Path
from flask import Flask, render_template, request, send_file, jsonify, session, redirect, url_for
import pandas as pd
import geopandas as gpd
from shapely.geometry import Polygon, MultiPolygon, GeometryCollection
from shapely.ops import unary_union
try:
    from shapely import make_valid
except ImportError:
    from shapely.validation import make_valid
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from functools import wraps
from werkzeug.security import check_password_hash

app = Flask(__name__)
app.secret_key = os.environ.get('ALRMS_SESSION_SECRET', 'alrms-local-session-v2-7-change-me')
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax')
LOGIN_USERNAME = os.environ.get('ALRMS_USERNAME', 'ALRMS')
DEFAULT_PASSWORD_HASH = 'pbkdf2:sha256:600000$alrms-v27-local$52d29729841c4eb54914bc5b24bea26e44050e1e7e77957f7e5238049d3ce4c5'
LOGIN_PASSWORD_HASH = os.environ.get('ALRMS_PASSWORD_HASH', DEFAULT_PASSWORD_HASH)

def login_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not session.get('logged_in'):
            if request.path.startswith(('/progress','/process','/inspect')):
                return jsonify({'error':'Login required.'}), 401
            return redirect(url_for('login'))
        return fn(*args, **kwargs)
    return wrapped


@app.after_request
def no_cache(response):
    response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
    response.headers['Pragma'] = 'no-cache'
    response.headers['Expires'] = '0'
    return response

app.config['MAX_CONTENT_LENGTH'] = 1024 * 1024 * 1024

WORK_ROOT = Path(tempfile.gettempdir()) / 'npaaad_clup_harmonizer'
WORK_ROOT.mkdir(exist_ok=True)

SAFDZ_ORDER = ['1','2','3','4','5','6','7','8','9','10','BU','WB','Others']
NPAAAD_ORDER = ['A','B','C','D','E','F','G','H','BU','WB','Others']
NPAAAD_AGR = {'A','B','C','D','E','G'}
SAFDZ_AGR = {'1','2','3','4','5','6','7'}

SAFDZ_DESC = {
 '1':'Strategic Crop Sub-Development Zone','2':'Strategic Livestock Sub-Development Zone','3':'Strategic Fishery Sub-Development Zone',
 '4':'Strategic Integrated Crop/Livestock Sub-Development Zone','5':'Strategic Integrated Crop/Fishery Sub-Development Zone',
 '6':'Strategic Integrated Crop/Livestock/Fishery Sub-Development Zone','7':'Strategic Integrated Fishery and Livestock Sub-Development Zone',
 '8':'Remaining NPAAAD','9':'Agro-Forestry Zone','10':'Forest/Watershed areas (critical watersheds including mangroves)',
 'BU':'Built-up areas (urban land, airport, roads and bridges)','Others':'Quarry, mine pit, barren land, rock land, river wash, landfill','WB':'Water Bodies'}
NPAAAD_DESC = {
 'A':'All irrigated lands/areas','B':'All irrigable lands already covered by irrigation projects with firm funding commitments',
 'C':'All alluvial plain lands highly suitable for agriculture, not irrigated',
 'D':'Agro-industrial croplands or lands presently planted to industrial crops that support existing agricultural infrastructure and agro-based enterprises',
 'E':'Highlands or areas at an elevation of 500 meters or above highly suitable for growing semi-temperate and high-value crops',
 'F':'Ecologically fragile agricultural lands where conversion may cause serious environmental degradation affecting mangroves and fish sanctuaries',
 'G':'All fishery areas as defined pursuant to Fisheries Code of 1998','H':'Forest/Watershed areas (critical watersheds including mangroves)',
 'BU':'Built-up areas (urban land, airport, roads and bridges)','Others':'Quarry, mine pit, barren land, rock land, river wash, beach sand, sand dunes, landfill'}


def extract_shapefile_zip(upload, dest: Path):
    dest.mkdir(parents=True, exist_ok=True)
    zpath = dest / upload.filename
    upload.save(zpath)
    with zipfile.ZipFile(zpath, 'r') as z:
        z.extractall(dest)
    shp_files = list(dest.rglob('*.shp'))
    if not shp_files:
        raise ValueError(f'No .shp found in {upload.filename}')
    return shp_files[0]


def inspect_uploaded_shapefile(upload, dest: Path):
    """Save a ZIP upload and return metadata used by the field-mapping UI."""
    shp = extract_shapefile_zip(upload, dest)
    gdf = gpd.read_file(shp)
    fields = [c for c in gdf.columns if c != 'geometry']
    crs_text = str(gdf.crs) if gdf.crs is not None else None
    epsg = gdf.crs.to_epsg() if gdf.crs is not None else None
    return gdf, {
        'filename': upload.filename,
        'crs': crs_text,
        'epsg': epsg,
        'authority': (':'.join(map(str, gdf.crs.to_authority())) if gdf.crs is not None and gdf.crs.to_authority() else None),
        'geometry_types': sorted([str(x) for x in gdf.geom_type.dropna().unique()]),
        'feature_count': int(len(gdf)),
        'fields': fields,
    }


def field_preview(gdf, field, limit=12):
    if not field or field not in gdf.columns:
        return []
    vals = gdf[field].dropna().astype(str).str.strip()
    vals = vals[vals != ''].drop_duplicates().head(limit).tolist()
    return vals


def suggest_field(fields, role):
    aliases = {
        'elu': ['ELU','EXISTING_LAND_USE','EXIST_LU','LANDUSE','LAND_USE','DESCRIPT','DESCRIPTION','CLASS'],
        'clup': ['CLUP','PROPOSED_LAND_USE','PROPOSED','PLU','LANDUSE','LAND_USE','DESCRIPT','DESCRIPTION','CLASS'],
        'npaaad': ['NPAAAD','NPAAAD_CL','NPAAAD_CLASS','CLASS','DESCRIPT','DESCRIPTION'],
        'safdz': ['SAFDZ','SAFDZ_CL','SAFDZ_CLASS','CLASS','DESCRIPT','DESCRIPTION'],
        'landcover': ['CLASS','LANDCOVER','LAND_COVER','LC_CLASS','LC_TYPE','DESCRIPT','DESCRIPTION'],
    }
    upper = {str(f).upper(): f for f in fields}
    for name in aliases.get(role, []):
        if name in upper:
            return upper[name]
    # Prefer text/category-looking fields rather than IDs/area fields.
    for f in fields:
        u = str(f).upper()
        if not any(x in u for x in ['OBJECTID','FID','SHAPE','AREA','LENGTH','PERIMETER']):
            return f
    return fields[0] if fields else None


def _polygon_part(geom):
    """Return polygonal content only.

    make_valid/clip/overlay can turn an invalid polygon into a GeometryCollection
    containing polygons, lines and points. GeoPandas overlay requires every input
    GeoDataFrame to contain one basic geometry family, so non-polygon remnants
    must be removed before the next overlay.
    """
    if geom is None or geom.is_empty:
        return None
    if isinstance(geom, (Polygon, MultiPolygon)):
        return geom
    if isinstance(geom, GeometryCollection):
        parts = []
        for part in geom.geoms:
            pg = _polygon_part(part)
            if pg is None or pg.is_empty:
                continue
            if isinstance(pg, Polygon):
                parts.append(pg)
            elif isinstance(pg, MultiPolygon):
                parts.extend(list(pg.geoms))
        if not parts:
            return None
        merged = unary_union(parts)
        return merged if isinstance(merged, (Polygon, MultiPolygon)) else None
    # Lines/points have no area and are not valid inputs for this polygon workflow.
    return None


def normalize_polygon_gdf(gdf: gpd.GeoDataFrame):
    """Make geometries valid and force a polygon-only GeoDataFrame."""
    gdf = gdf.copy()
    gdf = gdf[gdf.geometry.notnull()].copy()

    def _fix(g):
        if g is None or g.is_empty:
            return None
        try:
            g = make_valid(g) if not g.is_valid else g
        except Exception:
            pass
        return _polygon_part(g)

    gdf['geometry'] = gdf.geometry.apply(_fix)
    gdf = gdf[gdf.geometry.notnull() & ~gdf.geometry.is_empty].copy()
    # Explode MultiPolygons to simple Polygon records. This prevents mixed
    # Polygon/MultiPolygon exports and makes subsequent overlays more reliable.
    if not gdf.empty:
        gdf = gdf.explode(index_parts=False, ignore_index=True)
        gdf = gdf[gdf.geom_type == 'Polygon'].copy()
    return gdf


def clean_gdf(gdf: gpd.GeoDataFrame, keep_fields=None):
    gdf = normalize_polygon_gdf(gdf)
    if keep_fields:
        cols = [c for c in keep_fields if c in gdf.columns] + ['geometry']
        gdf = gdf[cols]
    return gdf


def safe_reproject(gdf, target_crs):
    if gdf.crs is None:
        raise ValueError('A layer has no CRS (.prj missing or unreadable).')
    # Avoid an expensive coordinate transformation when the layer already uses
    # the authoritative NPAAAD/SAFDZ CRS.
    if gdf.crs.equals(target_crs):
        return gdf
    return gdf.to_crs(target_crs)


def fast_prepare(gdf, field, target_crs):
    """Keep only the category + geometry, reproject once, repair only invalid rows."""
    if field not in gdf.columns:
        raise ValueError(f'Field "{field}" not found.')
    gdf = gdf[[field, 'geometry']].copy()
    gdf = gdf[gdf.geometry.notnull() & ~gdf.geometry.is_empty].copy()
    gdf[field] = gdf[field].astype(str).str.strip()
    gdf = safe_reproject(gdf, target_crs)
    # Critical performance change: valid polygons are left untouched.  v1.8
    # called Python/Shapely once for every feature, which made the 5% stage slow.
    valid = gdf.geometry.is_valid
    if not bool(valid.all()):
        bad = ~valid
        gdf.loc[bad, 'geometry'] = gdf.loc[bad, 'geometry'].apply(lambda g: _polygon_part(make_valid(g)))
        gdf = gdf[gdf.geometry.notnull() & ~gdf.geometry.is_empty].copy()
    return gdf


def fast_qa(gdf):
    """Cheap QA for live processing; expensive overlap scans are intentionally omitted."""
    try:
        invalid = int((~gdf.geometry.is_valid).sum())
    except Exception:
        invalid = -1
    return {'invalid': invalid, 'features': int(len(gdf))}


def detect_geometry_issues(gdf):
    invalid = int((~gdf.geometry.is_valid).sum())
    duplicates = int(gdf.geometry.apply(lambda g: g.wkb if g is not None else None).duplicated().sum())
    # Overlaps: spatial index based, unique pairs only
    overlaps = 0
    try:
        sindex = gdf.sindex
        seen = set()
        for i, geom in enumerate(gdf.geometry):
            if geom is None or geom.is_empty:
                continue
            for j in sindex.intersection(geom.bounds):
                if j <= i: continue
                key = (i,j)
                if key in seen: continue
                seen.add(key)
                try:
                    inter = geom.intersection(gdf.geometry.iloc[j])
                    if not inter.is_empty and inter.area > 0:
                        overlaps += 1
                except Exception:
                    pass
    except Exception:
        overlaps = -1
    holes = 0
    for geom in gdf.geometry:
        if isinstance(geom, Polygon):
            holes += len(geom.interiors)
        elif isinstance(geom, MultiPolygon):
            holes += sum(len(p.interiors) for p in geom.geoms)
    return {'invalid': invalid, 'duplicate_geometries': duplicates, 'overlap_pairs': overlaps, 'interior_holes': holes}


def dissolve_by_field(gdf, field):
    if field not in gdf.columns:
        raise ValueError(f'Field "{field}" not found.')
    gdf = gdf[[field, 'geometry']].copy()
    gdf[field] = gdf[field].astype(str).str.strip()
    out = gdf.dissolve(by=field, as_index=False)
    return normalize_polygon_gdf(out)




def fast_split_overlay(base, overlay, class_field):
    """Split base polygons only where an overlay actually intersects.

    Faster than a full union for this workflow because overlay is already clipped
    to the ELU boundary. We compute (1) intersections carrying the overlay class
    and (2) untouched base remainders outside the overlay coverage, then append.
    No outside-of-ELU union pieces are generated.
    """
    base = normalize_polygon_gdf(base)
    overlay = normalize_polygon_gdf(overlay[[class_field, 'geometry']].copy())
    if overlay.empty:
        out = base.copy()
        if class_field not in out.columns:
            out[class_field] = pd.NA
        return out

    # Spatial-indexed intersection creates only candidate pairs that overlap.
    inter = gpd.overlay(base, overlay, how='intersection', keep_geom_type=True)
    inter = normalize_polygon_gdf(inter)

    # Keep the portions of base that are not covered by this overlay. Because the
    # overlay was dissolved by classification first, union_all is very small.
    try:
        coverage = overlay.geometry.union_all()
    except AttributeError:
        coverage = overlay.geometry.unary_union
    rem = base.copy()
    rem.geometry = rem.geometry.difference(coverage)
    rem = normalize_polygon_gdf(rem)
    if class_field not in rem.columns:
        rem[class_field] = pd.NA

    if inter.empty:
        return rem
    out = pd.concat([inter, rem], ignore_index=True, sort=False)
    return normalize_polygon_gdf(gpd.GeoDataFrame(out, geometry='geometry', crs=base.crs))


def write_excel(summary_df, union_df, qa, out_path, poultry_by_elu=None):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    qa_ws = wb.create_sheet('QA_QC')
    attr_ws = wb.create_sheet('Union_Attribute_Table')

    title_fill = PatternFill('solid', fgColor='1F4E78')
    header_fill = PatternFill('solid', fgColor='D9EAF7')
    total_fill = PatternFill('solid', fgColor='E2F0D9')
    special_fill = PatternFill('solid', fgColor='FFF2CC')
    white_font = Font(color='FFFFFF', bold=True, size=12)
    bold_font = Font(bold=True)
    thin = Side(style='thin', color='B7B7B7')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def add_matrix(start_row, title, matrix, extra_cols, descriptions=None, poultry_by_elu=None):
        descriptions = descriptions or {}
        last_col = 1 + len(matrix.columns) + 1 + len(extra_cols)
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=last_col)
        c = ws.cell(start_row, 1, title); c.fill=title_fill; c.font=white_font; c.alignment=Alignment(horizontal='center')
        r = start_row + 1
        ws.cell(r,1,'ELU').font=bold_font
        for j,col in enumerate(matrix.columns,start=2): ws.cell(r,j,str(col)).font=bold_font
        gt_col=2+len(matrix.columns); ws.cell(r,gt_col,'Grand Total').font=bold_font
        for k,ex in enumerate(extra_cols,start=gt_col+1): ws.cell(r,k,ex).font=bold_font
        for j in range(1,last_col+1): ws.cell(r,j).fill=header_fill; ws.cell(r,j).border=border
        # Description row directly below the category codes.
        dr=r+1; ws.cell(dr,1,'Description').font=bold_font
        for j,col in enumerate(matrix.columns,start=2):
            ws.cell(dr,j,descriptions.get(str(col).strip().upper(),'')); ws.cell(dr,j).alignment=Alignment(wrap_text=True,vertical='top')
        ws.cell(dr,gt_col,'Area total (ha)')
        for k,ex in enumerate(extra_cols,start=gt_col+1):
            if ex=='Aglng-SZ': txt='SAFDZ 2 — Strategic Livestock Sub-Development Zone'
            elif ex=='PTA-SZ': txt='NPAAAD A, B, D, E, F, G'
            elif ex=='PDA-SZ': txt='NPAAAD C'
            elif ex=='Poultry Area (ha)': txt='Poultry shapefile area clipped within each ELU category'
            else: txt=''
            ws.cell(dr,k,txt); ws.cell(dr,k).alignment=Alignment(wrap_text=True,vertical='top')
        for j in range(1,last_col+1): ws.cell(dr,j).border=border
        ws.row_dimensions[dr].height=45
        rr=dr+1
        for idx,row in matrix.iterrows():
            ws.cell(rr,1,str(idx)); vals=[]
            for j,col in enumerate(matrix.columns,start=2):
                v=float(row[col]) if pd.notna(row[col]) else 0; ws.cell(rr,j,v); vals.append(v)
            ws.cell(rr,gt_col,sum(vals))
            for k,ex in enumerate(extra_cols,start=gt_col+1):
                if title.startswith('SAFDZ') and ex=='Aglng-SZ': v=float(row.get('2',0) or 0)
                elif title.startswith('NPAAAD') and ex=='PTA-SZ': v=sum(float(row.get(x,0) or 0) for x in ['A','B','D','E','F','G'])
                elif title.startswith('NPAAAD') and ex=='PDA-SZ': v=float(row.get('C',0) or 0)
                elif ex=='Poultry Area (ha)': v=float((poultry_by_elu or {}).get(str(idx),0) or 0)
                else: v=0
                ws.cell(rr,k,v)
            for j in range(1,last_col+1):
                ws.cell(rr,j).border=border
                if j>1: ws.cell(rr,j).number_format='#,##0.000'
            rr+=1
        ws.cell(rr,1,'Grand Total').font=bold_font
        for j,col in enumerate(matrix.columns,start=2): ws.cell(rr,j,float(matrix[col].sum()))
        ws.cell(rr,gt_col,float(matrix.to_numpy().sum()))
        for k,ex in enumerate(extra_cols,start=gt_col+1):
            if title.startswith('SAFDZ') and ex=='Aglng-SZ': v=float(matrix['2'].sum()) if '2' in matrix.columns else 0
            elif title.startswith('NPAAAD') and ex=='PTA-SZ': v=sum(float(matrix[x].sum()) for x in ['A','B','D','E','F','G'] if x in matrix.columns)
            elif title.startswith('NPAAAD') and ex=='PDA-SZ': v=float(matrix['C'].sum()) if 'C' in matrix.columns else 0
            elif ex=='Poultry Area (ha)': v=sum(float(x) for x in (poultry_by_elu or {}).values())
            else: v=0
            ws.cell(rr,k,v)
        for j in range(1,last_col+1):
            ws.cell(rr,j).fill=total_fill; ws.cell(rr,j).border=border
            if j>1: ws.cell(rr,j).number_format='#,##0.000'
        # Emphasize derived columns.
        for k,ex in enumerate(extra_cols,start=gt_col+1):
            for rownum in range(r,rr+1): ws.cell(rownum,k).fill=special_fill if rownum!=rr else total_fill
        return rr

    saf = summary_df['safdz']
    npaaad = summary_df['npaaad']
    end1 = add_matrix(1, 'SAFDZ — ELU Area Matrix (ha)', saf, ['Aglng-SZ'], SAFDZ_DESC)
    add_matrix(end1+3, 'NPAAAD — ELU Area Matrix (ha)', npaaad, ['PTA-SZ','PDA-SZ','Poultry Area (ha)'], NPAAAD_DESC, poultry_by_elu)
    ws.freeze_panes = 'B3'
    ws.column_dimensions['A'].width = 26
    for i in range(2, 24): ws.column_dimensions[get_column_letter(i)].width = 22

    # QA/QC sheet
    qa_ws.append(['Layer','Check','Value'])
    for c in qa_ws[1]: c.fill = header_fill; c.font = bold_font; c.border = border
    for layer,items in qa.items():
        for key,val in items.items():
            qa_ws.append([layer,key,val])
    for row in qa_ws.iter_rows():
        for c in row: c.border = border
    qa_ws.column_dimensions['A'].width=20; qa_ws.column_dimensions['B'].width=28; qa_ws.column_dimensions['C'].width=18

    # Attribute sheet
    df = union_df.drop(columns='geometry', errors='ignore').copy()
    for c in df.columns:
        if pd.api.types.is_datetime64_any_dtype(df[c]):
            df[c] = df[c].astype(str)
    attr_ws.append(list(df.columns))
    for c in attr_ws[1]: c.fill = header_fill; c.font = bold_font; c.border = border
    for row in df.itertuples(index=False, name=None):
        attr_ws.append([None if pd.isna(v) else v for v in row])
    for row in attr_ws.iter_rows():
        for c in row: c.border = border
    # Format hectare fields consistently.
    for idx, name in enumerate(df.columns, start=1):
        if str(name).upper().endswith('AREA_HA') or str(name).upper() == 'AREA_HA':
            for r in range(2, attr_ws.max_row + 1):
                attr_ws.cell(r, idx).number_format = '#,##0.000'
    attr_ws.freeze_panes='A2'

    wb.save(out_path)


@app.route('/login', methods=['GET', 'POST'])
def login():
    if session.get('logged_in'):
        return redirect(url_for('index'))
    error = None
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        if username == LOGIN_USERNAME and check_password_hash(LOGIN_PASSWORD_HASH, password):
            session.clear(); session['logged_in'] = True; session['username'] = username
            return redirect(url_for('index'))
        error = 'Invalid username or password.'
    return render_template('login.html', error=error)

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/')
@login_required
def index():
    return render_template('index.html')


@app.route('/health')
def health():
    return jsonify({'status': 'ok', 'app': 'NPAAAD/SAFDZ CLUP Harmonizer', 'version': '2.7'})


@app.route('/inspect', methods=['POST'])
@login_required
def inspect_layers():
    job = WORK_ROOT / ('inspect_' + next(tempfile._get_candidate_names()))
    job.mkdir(parents=True, exist_ok=True)
    roles = [
        ('elu_zip','elu',False), ('clup_zip','clup',False), ('npaaad_zip','npaaad',False),
        ('safdz_zip','safdz',False), ('landcover_zip','landcover',False), ('poultry_zip','poultry',False)
    ]
    result = {}
    try:
        for file_key, role, required in roles:
            up = request.files.get(file_key)
            if not up or not up.filename:
                if required:
                    raise ValueError(f'Missing required upload: {file_key}')
                continue
            gdf, meta = inspect_uploaded_shapefile(up, job/role)
            suggested = suggest_field(meta['fields'], role)
            meta['suggested_field'] = suggested
            meta['sample_values'] = field_preview(gdf, suggested)
            # ELU mapping UI needs the actual unique descriptions. Keep values
            # for each candidate field so changing the dropdown updates instantly.
            if role == 'elu':
                meta['field_values'] = {}
                for fld in meta['fields']:
                    vals = gdf[fld].dropna().astype(str).str.strip()
                    vals = vals[vals != ''].drop_duplicates().sort_values().head(1000).tolist()
                    meta['field_values'][fld] = vals
            result[role] = meta
        # NPAAAD + SAFDZ are the authoritative/master CRS. Other layers will
        # be reprojected to this CRS during processing.
        np_meta, saf_meta = result.get('npaaad'), result.get('safdz')
        master_crs = None
        master_ok = False
        master_error = None
        if np_meta and saf_meta:
            np_crs = gpd.read_file(list((job/'npaaad').rglob('*.shp'))[0]).crs
            saf_crs = gpd.read_file(list((job/'safdz').rglob('*.shp'))[0]).crs
            if np_crs is None or saf_crs is None:
                master_error = 'NPAAAD or SAFDZ has a missing/unreadable CRS.'
            elif np_crs.equals(saf_crs):
                master_ok = True
                auth = np_crs.to_authority()
                master_crs = ':'.join(map(str, auth)) if auth else np_crs.to_string()
            else:
                master_error = 'NPAAAD and SAFDZ use different CRS. They must match before processing.'
        result['_crs_summary'] = {
            'master_ok': master_ok,
            'master_crs': master_crs,
            'error': master_error
        }
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': str(e)}), 400
    finally:
        shutil.rmtree(job, ignore_errors=True)


JOBS = {}
JOBS_LOCK = threading.Lock()

def set_job(job_id, **kwargs):
    with JOBS_LOCK:
        JOBS.setdefault(job_id, {}).update(kwargs)

def extract_zip_path(zpath: Path, dest: Path):
    dest.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zpath, 'r') as z:
        z.extractall(dest)
    shp_files = list(dest.rglob('*.shp'))
    if not shp_files:
        raise ValueError(f'No .shp file found in {zpath.name}')
    return shp_files[0]

def run_background(job_id, job, uploads, fields, elu_mapping):
    started=time.time()
    def progress(pct, step):
        set_job(job_id, status='processing', percent=pct, step=step, elapsed=int(time.time()-started))
    try:
        progress(5, 'Reading uploaded shapefiles')
        shp_elu=extract_zip_path(uploads['elu_zip'], job/'elu')
        shp_np=extract_zip_path(uploads['npaaad_zip'], job/'npaaad')
        shp_saf=extract_zip_path(uploads['safdz_zip'], job/'safdz')
        elu_raw=gpd.read_file(shp_elu); np_raw=gpd.read_file(shp_np); saf_raw=gpd.read_file(shp_saf)
        if np_raw.crs is None or saf_raw.crs is None: raise ValueError('NPAAAD or SAFDZ has a missing/unreadable CRS (.prj).')
        if not np_raw.crs.equals(saf_raw.crs): raise ValueError('NPAAAD and SAFDZ have different CRS. Processing stopped to protect spatial alignment.')
        target_crs=np_raw.crs
        # v1.8: the former full overlap/duplicate QA scan was the main reason the
        # UI appeared stuck at 5%. Use a fast validity/count QA during the run.
        qa={'ELU_before':fast_qa(elu_raw),'NPAAAD_before':fast_qa(np_raw),'SAFDZ_before':fast_qa(saf_raw)}
        progress(15, 'Preparing layers')
        # v1.9: copy user-selected fields into unique internal names BEFORE any
        # overlay. GeoPandas otherwise suffixes same-named fields (_1/_2), which
        # caused the v1.8 KeyError: 'SAFDZ' during summary generation.
        internal={'elu':'ELU_CLS','clup':'CLUP_CLS','npaaad':'NP_CLS','safdz':'SZ_CLS','landcover':'LC_CLS'}
        poultry_clip=None; poultry_by_elu={}
        for label, gdf, key in [('ELU',elu_raw,'elu'),('NPAAAD',np_raw,'npaaad'),('SAFDZ',saf_raw,'safdz')]:
            if fields[key] not in gdf.columns:
                raise ValueError(f"{label} classification field '{fields[key]}' was not found. Please select a valid field before processing.")
        elu_p=fast_prepare(elu_raw,fields['elu'],target_crs).rename(columns={fields['elu']:internal['elu']})
        np_p=fast_prepare(np_raw,fields['npaaad'],target_crs).rename(columns={fields['npaaad']:internal['npaaad']})
        saf_p=fast_prepare(saf_raw,fields['safdz'],target_crs).rename(columns={fields['safdz']:internal['safdz']})
        progress(27, 'Preparing ELU boundary')
        elu=dissolve_by_field(elu_p,internal['elu'])
        try: boundary_geom=elu.geometry.union_all()
        except Exception: boundary_geom=elu.geometry.unary_union
        elu_boundary=gpd.GeoDataFrame(geometry=[boundary_geom],crs=elu.crs)
        # v2.4 speed optimization: spatially clip first, dissolve second. This
        # avoids dissolving thousands of polygons that are outside the ELU.
        progress(36, 'Filtering NPAAAD and SAFDZ to ELU')
        np_local=normalize_polygon_gdf(gpd.clip(np_p,elu_boundary,keep_geom_type=True))
        saf_local=normalize_polygon_gdf(gpd.clip(saf_p,elu_boundary,keep_geom_type=True))
        np_clip=dissolve_by_field(np_local,internal['npaaad'])
        saf_clip=dissolve_by_field(saf_local,internal['safdz'])
        progress(48, 'Clipping NPAAAD and SAFDZ to ELU')
        # v2.2 FAST OVERLAY ENGINE: avoid full polygon union. NPAAAD/SAFDZ are
        # already clipped to ELU and dissolved by class, so split only polygons
        # that actually intersect and append the untouched ELU remainder.
        progress(48, 'Intersecting relevant NPAAAD polygons')
        union1 = fast_split_overlay(elu, np_clip, internal['npaaad'])
        if internal['elu'] not in union1.columns:
            raise ValueError('ELU classification was lost during NPAAAD intersection.')

        progress(62, 'Intersecting relevant SAFDZ polygons')
        union2 = fast_split_overlay(union1, saf_clip, internal['safdz'])
        for f in (internal['elu'], internal['npaaad'], internal['safdz']):
            if f not in union2.columns:
                raise ValueError(f'Classification field was lost after fast overlay: {f}')
        if 'landcover_zip' in uploads:
            progress(70, 'Processing Land Cover Map 2025')
            shp_lc=extract_zip_path(uploads['landcover_zip'],job/'landcover'); lc_raw=gpd.read_file(shp_lc)
            qa['LandCover_before']=fast_qa(lc_raw)
            if fields['landcover'] not in lc_raw.columns:
                raise ValueError(f"Land Cover classification field '{fields['landcover']}' was not found. Please select a valid field before processing.")
            lc_p=fast_prepare(lc_raw,fields['landcover'],target_crs).rename(columns={fields['landcover']:internal['landcover']})
            # Clip before dissolve: only Land Cover polygons touching ELU are processed.
            lc_local=normalize_polygon_gdf(gpd.clip(lc_p,elu_boundary,keep_geom_type=True))
            lc_clip=dissolve_by_field(lc_local,internal['landcover'])
            progress(72, 'Intersecting relevant Land Cover polygons')
            union2 = fast_split_overlay(union2, lc_clip, internal['landcover'])
            if internal['landcover'] not in union2.columns:
                union2[internal['landcover']] = pd.NA
        # Approved CLUP is optional. When supplied, split the harmonized layer by
        # CLUP so the approved-CLUP decision matrix can be evaluated per polygon.
        if 'clup_zip' in uploads:
            progress(76, 'Intersecting Approved CLUP polygons')
            shp_clup=extract_zip_path(uploads['clup_zip'],job/'clup'); clup_raw=gpd.read_file(shp_clup)
            qa['CLUP_before']=fast_qa(clup_raw)
            if fields['clup'] not in clup_raw.columns:
                raise ValueError(f"CLUP classification field '{fields['clup']}' was not found. Please select a valid field before processing.")
            clup_p=fast_prepare(clup_raw,fields['clup'],target_crs).rename(columns={fields['clup']:internal['clup']})
            clup_local=normalize_polygon_gdf(gpd.clip(clup_p,elu_boundary,keep_geom_type=True))
            clup_clip=dissolve_by_field(clup_local,internal['clup'])
            union2=fast_split_overlay(union2,clup_clip,internal['clup'])
        if 'poultry_zip' in uploads:
            progress(78, 'Clipping Poultry layer to ELU')
            shp_poultry=extract_zip_path(uploads['poultry_zip'],job/'poultry'); poultry_raw=gpd.read_file(shp_poultry)
            qa['Poultry_before']=fast_qa(poultry_raw)
            if poultry_raw.crs is None: raise ValueError('Poultry layer has a missing/unreadable CRS (.prj).')
            poultry_p=safe_reproject(poultry_raw[['geometry']].copy(),target_crs)
            poultry_p=normalize_polygon_gdf(poultry_p)
            poultry_clip=normalize_polygon_gdf(gpd.clip(poultry_p,elu_boundary,keep_geom_type=True))
            if not poultry_clip.empty:
                try: poultry_geom=poultry_clip.geometry.union_all()
                except Exception: poultry_geom=poultry_clip.geometry.unary_union
                # Poultry hectares per ELU without splitting the harmonized geometry.
                for _,er in elu.iterrows():
                    try: poultry_by_elu[str(er[internal['elu']])]=round(er.geometry.intersection(poultry_geom).area/10000.0,3)
                    except Exception: poultry_by_elu[str(er[internal['elu']])]=0.0
                union2['POULTRY_HA']=union2.geometry.apply(lambda g: round(g.intersection(poultry_geom).area/10000.0,3) if g is not None and not g.is_empty else 0.0)
            else:
                union2['POULTRY_HA']=0.0
        progress(80, 'Computing area in hectares')
        union2=normalize_polygon_gdf(union2)

        def norm(v):
            if pd.isna(v): return ''
            return str(v).strip()
        def broad_lu(v):
            x=norm(v).upper().replace('_',' ').replace('-',' ').strip()
            if x in {'BU','BUILT UP','BUILTUP'} or 'BUILT UP' in x: return 'BU'
            if x in {'AGR','AGRICULTURE','AGRICULTURAL'} or x.startswith('AGR '): return 'AGR'
            return x
        def satellite(v):
            x=norm(v).upper().replace('_',' ').replace('-',' ').strip()
            if x in {'ANNUAL CROP','ANNUAL CROPS','PERENNIAL CROP','PERENNIAL CROPS'}: return 'AGR'
            if x in {'BUILT UP','BUILTUP','BU'}: return 'BU'
            return ''
        def np_group(v):
            x=norm(v).upper()
            if x in NPAAAD_AGR: return 'AGR'
            if x == 'BU': return 'BU'
            return x
        def sz_group(v):
            x=norm(v).upper()
            if x in SAFDZ_AGR: return 'AGR'
            if x == 'BU': return 'BU'
            return x
        def ag_context(npv, szv):
            # SAFDZ is the more specific agricultural zoning when codes 1-7 exist.
            s=norm(szv).upper(); n=norm(npv).upper()
            if s in SAFDZ_AGR: return ('AGR', s)
            if n in NPAAAD_AGR: return ('AGR', n)
            if s == 'BU' or n == 'BU': return ('BU', 'BU')
            return ('', '')
        def decide(row):
            sat=satellite(row.get(internal['landcover'], ''))
            zone, code=ag_context(row.get(internal['npaaad'],''), row.get(internal['safdz'],''))
            if internal['clup'] in row.index and norm(row.get(internal['clup'],'')):
                lu=broad_lu(row.get(internal['clup'],''))
                if lu=='BU' and zone=='AGR' and sat=='BU': return 'BU'
                if lu=='BU' and zone=='AGR' and sat=='AGR': return f'{code}/BU' if code else '?/BU'
                if lu=='AGR' and zone=='BU' and sat=='BU': return 'AGR'
            else:
                raw_elu=norm(row.get(internal['elu'],''))
                lu=elu_mapping.get(raw_elu, broad_lu(raw_elu))
                if lu=='OTHER': return ''
                if lu=='BU' and zone=='AGR' and sat=='BU': return 'BU'
                if lu=='BU' and zone=='AGR' and sat=='AGR': return 'do not harmonize'
                if lu=='AGR' and zone=='BU' and sat=='BU': return 'Ask LGU'
            return ''

        union2['Decision']=union2.apply(decide,axis=1)
        union2['AREA_HA']=(union2.geometry.area/10000.0).round(3)
        u=union2.copy()
        required=[internal['elu'],internal['npaaad'],internal['safdz']]
        missing=[f for f in required if f not in u.columns]
        if missing:
            raise ValueError('Required classification fields were lost during overlay: ' + ', '.join(missing))
        for f in required:
            u[f]=u[f].where(u[f].notna(), '').astype(str).str.strip()
        saf_matrix=pd.pivot_table(u,values='AREA_HA',index=internal['elu'],columns=internal['safdz'],aggfunc='sum',fill_value=0)
        np_matrix=pd.pivot_table(u,values='AREA_HA',index=internal['elu'],columns=internal['npaaad'],aggfunc='sum',fill_value=0)
        saf_cols=[c for c in SAFDZ_ORDER if c in saf_matrix.columns]+[c for c in saf_matrix.columns if c not in SAFDZ_ORDER]
        np_cols=[c for c in NPAAAD_ORDER if c in np_matrix.columns]+[c for c in np_matrix.columns if c not in NPAAAD_ORDER]
        saf_matrix=saf_matrix.reindex(columns=saf_cols); np_matrix=np_matrix.reindex(columns=np_cols)
        qa['Final']={'crs':str(union2.crs),'features':len(union2),'total_area_ha':round(float(union2['AREA_HA'].sum()),3),'geometry_types':', '.join(sorted(union2.geom_type.unique()))}
        progress(90, 'Generating Excel, Shapefile and QA/QC report')
        out_dir=job/'outputs'; out_dir.mkdir(); shp_dir=out_dir/'union_shapefile'; shp_dir.mkdir()
        shp_path=shp_dir/'HARMONIZED_UNION.shp'; union2.to_file(shp_path)
        shp_zip=out_dir/'HARMONIZED_UNION.zip'
        with zipfile.ZipFile(shp_zip,'w',zipfile.ZIP_DEFLATED) as z:
            for q in shp_dir.iterdir(): z.write(q,q.name)
        xlsx=out_dir/'HARMONIZATION_SUMMARY.xlsx'; write_excel({'safdz':saf_matrix,'npaaad':np_matrix},union2,qa,xlsx,poultry_by_elu)
        bundle=out_dir/'HARMONIZATION_RESULTS.zip'
        with zipfile.ZipFile(bundle,'w',zipfile.ZIP_DEFLATED) as z:
            z.write(xlsx,xlsx.name); z.write(shp_zip,shp_zip.name)
            if poultry_clip is not None:
                pdir=out_dir/'poultry_clipped'; pdir.mkdir(exist_ok=True); pshp=pdir/'POULTRY_CLIPPED.shp'; poultry_clip.to_file(pshp)
                pzip=out_dir/'POULTRY_CLIPPED.zip'
                with zipfile.ZipFile(pzip,'w',zipfile.ZIP_DEFLATED) as pz:
                    for q in pdir.iterdir(): pz.write(q,q.name)
                z.write(pzip,pzip.name)
            qa_json=out_dir/'QA_QC.json'; qa_json.write_text(json.dumps(qa,indent=2),encoding='utf-8'); z.write(qa_json,qa_json.name)
        set_job(job_id,status='complete',percent=100,step='Processing Complete',elapsed=int(time.time()-started),bundle=str(bundle))
    except Exception as e:
        set_job(job_id,status='error',percent=JOBS.get(job_id,{}).get('percent',0),step='Processing Failed',elapsed=int(time.time()-started),error=str(e),details=traceback.format_exc())

@app.route('/process_start', methods=['POST'])
@login_required
def process_start():
    try:
        for r in ['elu_zip','npaaad_zip','safdz_zip']:
            if r not in request.files or not request.files[r].filename: raise ValueError(f'Missing required upload: {r}')
        job_id=uuid.uuid4().hex; job=WORK_ROOT/job_id; up=job/'uploads'; up.mkdir(parents=True)
        uploads={}
        for name in ['elu_zip','npaaad_zip','safdz_zip','landcover_zip','clup_zip','poultry_zip']:
            f=request.files.get(name)
            if f and f.filename:
                path=up/(name+'.zip'); f.save(path); uploads[name]=path
        fields={'elu':request.form.get('elu_field','ELU').strip(),'clup':request.form.get('clup_field','CLUP').strip(),'npaaad':request.form.get('npaaad_field','NPAAAD').strip(),'safdz':request.form.get('safdz_field','SAFDZ').strip(),'landcover':request.form.get('landcover_field','CLASS').strip()}
        try:
            elu_mapping=json.loads(request.form.get('elu_mapping','{}') or '{}')
        except Exception:
            elu_mapping={}
        elu_mapping={str(k).strip():str(v).strip().upper() for k,v in elu_mapping.items() if str(v).strip().upper() in {'AGR','BU','OTHER'}}
        set_job(job_id,status='queued',percent=0,step='Preparing data',elapsed=0)
        threading.Thread(target=run_background,args=(job_id,job,uploads,fields,elu_mapping),daemon=True).start()
        return jsonify({'job_id':job_id})
    except Exception as e: return jsonify({'error':str(e)}),400

@app.route('/progress/<job_id>')
@login_required
def job_progress(job_id):
    with JOBS_LOCK: data=dict(JOBS.get(job_id,{}))
    if not data: return jsonify({'error':'Job not found'}),404
    data.pop('bundle',None); return jsonify(data)

@app.route('/download/<job_id>')
@login_required
def job_download(job_id):
    with JOBS_LOCK: data=dict(JOBS.get(job_id,{}))
    if data.get('status')!='complete' or not data.get('bundle'): return jsonify({'error':'Results are not ready.'}),409
    return send_file(data['bundle'],as_attachment=True,download_name='HARMONIZATION_RESULTS.zip')

if __name__ == '__main__':
    port=int(os.environ.get('PORT','5055')); app.run(host='127.0.0.1',port=port,debug=False,use_reloader=False,threaded=True)
